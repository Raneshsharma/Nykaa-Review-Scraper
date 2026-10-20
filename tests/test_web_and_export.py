import time

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from conftest import fast_scraper
from nykaa_scraper.exporter import export_to_csv, export_to_excel, export_to_json


@pytest.fixture
def client(mock_api, monkeypatch, tmp_path):
    mock_api(n_reviews=250)
    import app as webapp
    monkeypatch.setattr(webapp, "OUTPUT_DIR", tmp_path)
    return TestClient(webapp.app)


def wait_for(client, job_id, timeout=60):
    end = time.time() + timeout
    while time.time() < end:
        job = client.get(f"/api/jobs/{job_id}").json()
        if job["status"] != "running":
            return job
        time.sleep(0.2)
    raise AssertionError("job did not finish")


def test_web_job_flow_and_downloads(client):
    assert client.get("/").status_code == 200
    start = client.post("/api/scrape", json={"target": "123456", "limit": 0}).json()
    job = wait_for(client, start["job_id"])
    assert job["status"] == "done", job.get("error")
    assert len(job["data"]["reviews"]) == 250
    assert job["data"]["summary"]["total_available"] == 250
    for fmt in ("csv", "xlsx", "json"):
        r = client.get(f"/api/download/{job['result_id']}/{fmt}")
        assert r.status_code == 200 and len(r.content) > 100


def test_web_job_error_is_reported(client):
    start = client.post("/api/scrape", json={"target": "https://nykaa.app.link/none"}).json()
    job = wait_for(client, start["job_id"])
    assert job["status"] == "error" and job["error"]


def test_unknown_job_404(client):
    assert client.get("/api/jobs/nope").status_code == 404


def test_exports_with_awkward_text(mock_api, tmp_path):
    mock_api(n_reviews=80)
    result = fast_scraper().scrape("123456")
    result.reviews[0].description = "emoji 💄 and control \x01 chars"  # exporter must not crash
    export_to_csv(result, str(tmp_path / "r.csv"))
    export_to_json(result, str(tmp_path / "r.json"))
    export_to_excel(result, str(tmp_path / "r.xlsx"))
    wb = load_workbook(tmp_path / "r.xlsx")
    assert wb["Reviews"].max_row == 81
