"""
Export utilities for saving reviews to CSV, Excel, and JSON formats.
"""

import json
from pathlib import Path
from typing import List, Dict, Any
import pandas as pd

from .models import ScrapeResult, Review


def reviews_to_dataframe(reviews: List[Review]) -> pd.DataFrame:
    """Converts a list of Review objects to a clean pandas DataFrame with only requested fields."""
    records = []
    for r in reviews:
        records.append({
            "Reviewer": r.reviewer_name,
            "Rating": r.rating,
            "Title": r.title,
            "Written Text": r.description,
        })
    return pd.DataFrame(records)


def export_to_csv(result: ScrapeResult, file_path: str) -> str:
    """
    Exports scraped reviews to a CSV file.
    Uses 'utf-8-sig' to support emojis and special characters directly in Excel.
    """
    Path(file_path).parent.mkdir(parents=True, exist_ok=True)
    df = reviews_to_dataframe(result.reviews)
    df.to_csv(file_path, index=False, encoding="utf-8-sig")
    return file_path


def export_to_excel(result: ScrapeResult, file_path: str) -> str:
    """
    Exports scraped reviews to a beautifully formatted Excel (.xlsx) file.
    Includes a 'Reviews' data sheet and a 'Product Summary' sheet.
    """
    Path(file_path).parent.mkdir(parents=True, exist_ok=True)
    df_reviews = reviews_to_dataframe(result.reviews)

    # Summary table
    summary_data = [
        {"Attribute": "Product ID", "Value": result.product.id},
        {"Attribute": "Product Name", "Value": result.product.name},
        {"Attribute": "Brand", "Value": result.product.brand},
        {"Attribute": "Price (INR)", "Value": f"₹{result.product.price}" if result.product.price else "N/A"},
        {"Attribute": "Nykaa URL", "Value": result.product.url},
        {"Attribute": "Total Reviews on Nykaa", "Value": result.product.total_reviews_count},
        {"Attribute": "Total Scraped Reviews", "Value": result.summary.total_scraped},
        {"Attribute": "Average Rating (Scraped)", "Value": f"{result.summary.average_rating} / 5.0"},
        {"Attribute": "5-Star Reviews", "Value": result.summary.star_distribution.get("5", 0)},
        {"Attribute": "4-Star Reviews", "Value": result.summary.star_distribution.get("4", 0)},
        {"Attribute": "3-Star Reviews", "Value": result.summary.star_distribution.get("3", 0)},
        {"Attribute": "2-Star Reviews", "Value": result.summary.star_distribution.get("2", 0)},
        {"Attribute": "1-Star Reviews", "Value": result.summary.star_distribution.get("1", 0)},
        {"Attribute": "Verified Buyers", "Value": result.summary.verified_buyer_count},
        {"Attribute": "Reviews With Images", "Value": result.summary.with_images_count},
        {"Attribute": "Scraped At", "Value": result.scraped_at},
    ]
    df_summary = pd.DataFrame(summary_data)

    with pd.ExcelWriter(file_path, engine="openpyxl") as writer:
        df_reviews.to_excel(writer, sheet_name="Reviews", index=False)
        df_summary.to_excel(writer, sheet_name="Product Summary", index=False)

        # Style the sheets with openpyxl
        try:
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            from openpyxl.utils import get_column_letter

            wb = writer.book

            header_fill = PatternFill(start_color="FC2779", end_color="FC2779", fill_type="solid")
            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            thin_border = Border(
                left=Side(style='thin', color='DDDDDD'),
                right=Side(style='thin', color='DDDDDD'),
                top=Side(style='thin', color='DDDDDD'),
                bottom=Side(style='thin', color='DDDDDD'),
            )

            for ws in wb.worksheets:
                ws.views.sheetView[0].showGridLines = True
                for cell in ws[1]:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

                for row in ws.iter_rows(min_row=2):
                    for cell in row:
                        cell.border = thin_border
                        cell.alignment = Alignment(vertical="center")

                # Auto-adjust column widths
                for col in ws.columns:
                    max_len = 0
                    col_letter = get_column_letter(col[0].column)
                    for cell in col:
                        val_str = str(cell.value or "")
                        if len(val_str) > max_len:
                            max_len = len(val_str)
                    ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 60)

        except Exception:
            pass  # Fallback to standard excel without custom styles

    return file_path


def export_to_json(result: ScrapeResult, file_path: str) -> str:
    """Exports scraped reviews and summary to a structured JSON file with requested review fields."""
    Path(file_path).parent.mkdir(parents=True, exist_ok=True)
    clean_reviews = [
        {
            "reviewer": r.reviewer_name,
            "rating": r.rating,
            "title": r.title,
            "written_text": r.description,
        }
        for r in result.reviews
    ]
    data = {
        "product": result.product.name,
        "product_id": result.product.id,
        "average_rating": result.summary.average_rating,
        "total_scraped": result.summary.total_scraped,
        "reviews": clean_reviews,
        "scraped_at": result.scraped_at,
    }
    with open(file_path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)
    return file_path
